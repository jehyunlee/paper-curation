#!/usr/bin/env python3
"""AI4Science 발표 덱 → PowerPoint(.pptx).

`build_slide_deck.build_slides()` 가 만든 슬라이드(불릿판)를 그대로 16:9 PPTX 로
옮기고, 같은 장의 줄글(`build_slide_essay.build()`)을 **발표자 노트**에 넣는다.
데이터·집필본·레퍼런스 선정은 두 빌더의 것을 재사용하므로 HTML/MD 와 내용이 같다.

  - `reports/build/{topic}_slides_<N>.pptx`

레퍼런스 하이퍼링크는 `--link-base`(기본 `../../docs/papers`) 를 그대로 쓴다 —
PPTX 파일 위치 기준 상대 경로라 `reports/build/` 에 두고 열면 리뷰 문서로 이어진다.
배포 사이트로 걸려면 `--link-base https://<host>/papers` 처럼 절대 URL 을 준다.

Usage:
  PYTHONUTF8=1 python pipeline/build_slide_pptx.py --topic ai4s
"""
import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PIPE = Path(__file__).resolve().parent
sys.path.insert(0, str(PIPE))
PROJECT_ROOT = PIPE.parent

import build_slide_deck as B  # noqa: E402
import build_slide_essay as E  # noqa: E402

REPORTS = PROJECT_ROOT / "reports"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)
INNER_W = W - 2 * MARGIN
FONT_LATIN = "Helvetica Neue"
FONT_EA = "Malgun Gothic"  # Office(Win·Mac) 동봉; 없으면 시스템 고딕으로 대체

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
LINE = RGBColor(0xDD, 0xDD, 0xDD)
PANEL = RGBColor(0xF6, 0xF6, 0xF6)
LINK = RGBColor(0x1F, 0x4E, 0x9A)

_MD_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


def _set_font(run, size, bold=False, italic=False, color=INK, name=FONT_LATIN):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = name
    f.color.rgb = color
    # 동아시아 글꼴을 비워 두거나 시스템 전용 글꼴(Apple SD Gothic Neo)을 주면
    # PowerPoint 가 한글을 바탕(Batang)으로 대체한다. Office 가 직접 동봉하는 맑은 고딕을 쓴다.
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", FONT_EA)


def _add_runs(paragraph, text, size, color=INK, bold=False, italic=False):
    """`**굵게**` 마크다운을 run 으로 풀어 넣는다."""
    pos = 0
    for m in _MD_BOLD.finditer(text):
        if m.start() > pos:
            _set_font(paragraph.add_run(), size, bold, italic, color)
            paragraph.runs[-1].text = text[pos:m.start()]
        _set_font(paragraph.add_run(), size, True, italic, color)
        paragraph.runs[-1].text = m.group(1)
        pos = m.end()
    if pos < len(text):
        _set_font(paragraph.add_run(), size, bold, italic, color)
        paragraph.runs[-1].text = text[pos:]


def _textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, fill=None, line=None, inset=0.08):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(inset)
    tf.margin_top = tf.margin_bottom = Inches(inset * 0.6)
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is not None:
        box.line.color.rgb = line
        box.line.width = Pt(0.75)
    return box, tf


def _rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _para(tf, first):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def _body_font_size(points, has_table):
    chars = sum(len(p) for p in points)
    if has_table:
        return 11 if chars > 700 else 12
    if chars > 1100:
        return 11
    if chars > 800:
        return 12
    if chars > 500:
        return 13
    return 14


def _add_table(slide, table, x, y, w, h):
    head, rows = table["head"], table["rows"]
    n_rows, n_cols = len(rows) + 1, len(head)
    row_h = Emu(int(min(Inches(0.3), h / n_rows)))
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, row_h * n_rows)
    tbl = shape.table
    # 기본 테마(파란 머리글·줄무늬)를 끄고 선만 남긴다
    tbl_pr = tbl._tbl.tblPr
    tbl_pr.set("firstRow", "0")
    tbl_pr.set("bandRow", "0")
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is not None:
        style_id.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid

    # 열 폭: 각 열의 최장 문자열 길이에 비례(한글은 2배 가중), 최소 폭 보장
    def _wlen(t):
        return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(t))
    longest = [max([_wlen(head[j])] + [_wlen(r[j]) for r in rows]) for j in range(n_cols)]
    weights = [max(6, min(v, 34)) for v in longest]
    total = sum(weights)
    for j, col in enumerate(tbl.columns):
        col.width = Emu(int(w * weights[j] / total))

    def _fill(cell, text, size, color=INK, bold=False):
        cell.text = ""
        cell.margin_left = cell.margin_right = Inches(0.05)
        cell.margin_top = cell.margin_bottom = Inches(0.015)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.background()
        _add_runs(cell.text_frame.paragraphs[0], text, size, color, bold=bold)

    size = 9 if n_rows > 8 else 10
    for j, cell_text in enumerate(head):
        _fill(tbl.cell(0, j), str(cell_text), size, MUTED, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, cell_text in enumerate(row):
            txt = str(cell_text)
            # 코퍼스 연도표의 '█' 막대는 PPT 에서 폭이 들쭉날쭉하니 길이만 남긴다
            if txt and set(txt) <= {"█"}:
                txt = "▇" * max(1, len(txt) // 2)
            _fill(tbl.cell(i, j), txt, size)
    for r in tbl.rows:
        r.height = row_h
    return shape, row_h * n_rows


def add_slide(prs, s, accent, notes_text):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)

    # 상단: 번호 태그 + 종류 + 파트
    _rect(slide, MARGIN, Inches(0.32), Inches(0.62), Inches(0.28), INK)
    _, tf = _textbox(slide, MARGIN, Inches(0.30), Inches(0.62), Inches(0.32), anchor=MSO_ANCHOR.MIDDLE, inset=0.02)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _add_runs(p, f"S{s['no']:02d}", 10, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    _, tf = _textbox(slide, MARGIN + Inches(0.7), Inches(0.30), Inches(3), Inches(0.32), anchor=MSO_ANCHOR.MIDDLE, inset=0.02)
    _add_runs(tf.paragraphs[0], s.get("kind", ""), 10, accent, bold=True)
    _, tf = _textbox(slide, W - MARGIN - Inches(6), Inches(0.30), Inches(6), Inches(0.32), anchor=MSO_ANCHOR.MIDDLE, inset=0.02)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _add_runs(tf.paragraphs[0], s.get("part", ""), 10, MUTED)

    # 제목
    y = Inches(0.68)
    _, tf = _textbox(slide, MARGIN, y, INNER_W, Inches(0.7), anchor=MSO_ANCHOR.TOP, inset=0.02)
    _add_runs(tf.paragraphs[0], s.get("title_ko", ""), 26, bold=True)
    y += Inches(0.56)
    if s.get("title_en"):
        _, tf = _textbox(slide, MARGIN, y, INNER_W, Inches(0.3), inset=0.02)
        _add_runs(tf.paragraphs[0], s["title_en"], 12, MUTED, italic=True)
        y += Inches(0.3)

    # 배지 · 카테고리 메타
    meta_bits = []
    if s.get("badges"):
        meta_bits.append("   ".join(f"[{b}]" for b in s["badges"]))
    if s.get("cat_meta"):
        meta_bits.append(s["cat_meta"])
    if meta_bits:
        _, tf = _textbox(slide, MARGIN, y, INNER_W, Inches(0.3), inset=0.02)
        _add_runs(tf.paragraphs[0], " · ".join(meta_bits), 10, MUTED)
        y += Inches(0.3)

    # 파트 도입(context)
    if s.get("context"):
        _, tf = _textbox(slide, MARGIN, y, INNER_W, Inches(0.36), anchor=MSO_ANCHOR.MIDDLE, fill=PANEL, inset=0.1)
        _add_runs(tf.paragraphs[0], s["context"], 11, MUTED)
        y += Inches(0.42)

    # 헤드라인
    if s.get("headline"):
        _rect(slide, MARGIN, y + Inches(0.04), Inches(0.06), Inches(0.5), accent)
        _, tf = _textbox(slide, MARGIN + Inches(0.14), y, INNER_W - Inches(0.14), Inches(0.58), anchor=MSO_ANCHOR.MIDDLE, inset=0.04)
        _add_runs(tf.paragraphs[0], s["headline"], 15, bold=True)
        y += Inches(0.64)

    # 본문: 표(있으면 왼쪽) + 불릿
    refs = s.get("refs") or []
    sowhat_h = Inches(0.62) if s.get("sowhat") else 0
    refs_h = Inches(0.22) * min(len(refs), 4) + (Inches(0.1) if refs else 0)
    body_bottom = H - MARGIN - sowhat_h - refs_h - Inches(0.08)
    body_h = max(Inches(1.2), body_bottom - y)
    points = list(s.get("points") or [])
    extra = s.get("extra") or {}
    table = s.get("table")
    text_x, text_w = MARGIN, INNER_W
    if table:
        tw = INNER_W * (0.58 if len(table["head"]) >= 4 else 0.36)
        _add_table(slide, table, MARGIN, y, tw, body_h)
        text_x, text_w = MARGIN + tw + Inches(0.25), INNER_W - tw - Inches(0.25)
    size = _body_font_size(points + [v for vs in extra.values() for v in vs], bool(table))
    _, tf = _textbox(slide, text_x, y, text_w, body_h, inset=0.04)
    first = True
    for pt in points:
        p = _para(tf, first)
        first = False
        p.space_after = Pt(4)
        if pt.startswith("**") and pt.endswith("—**"):
            _add_runs(p, pt, size, accent)
            continue
        _add_runs(p, "•  " + pt, size)
    for label, items in extra.items():
        if not items:
            continue
        p = _para(tf, first)
        first = False
        p.space_before = Pt(6)
        _add_runs(p, label, size - 1, accent, bold=True)
        for it in items:
            p = tf.add_paragraph()
            p.space_after = Pt(2)
            _add_runs(p, "–  " + str(it), size - 1, MUTED)
    if s.get("tools"):
        p = _para(tf, first)
        first = False
        p.space_before = Pt(6)
        _add_runs(p, "대표 도구 · 시스템: " + " · ".join(s["tools"]), 9, MUTED)

    # 레퍼런스 (하이퍼링크)
    y_refs = H - MARGIN - sowhat_h - refs_h
    if refs:
        _rect(slide, MARGIN, y_refs, INNER_W, Emu(9525), LINE)
        _, tf = _textbox(slide, MARGIN, y_refs + Inches(0.04), INNER_W, refs_h, inset=0.02)
        first = True
        for i, r in enumerate(refs[:4], start=1):
            p = _para(tf, first)
            first = False
            p.space_after = Pt(1)
            run = p.add_run()
            run.text = f"[{i}] {r['title']}"
            _set_font(run, 9, color=LINK)
            if r.get("url"):
                run.hyperlink.address = r["url"]
            tail = p.add_run()
            meta = " — " + ", ".join(x for x in (r.get("author"), r.get("date")) if x)
            if r.get("citations"):
                meta += f" · 인용 {r['citations']:,}"
            tail.text = meta
            _set_font(tail, 9, color=MUTED)

    # 우리에게 무엇인가
    if s.get("sowhat"):
        y_so = H - MARGIN - sowhat_h + Inches(0.06)
        _, tf = _textbox(slide, MARGIN, y_so, INNER_W, sowhat_h - Inches(0.06), anchor=MSO_ANCHOR.MIDDLE,
                         fill=PANEL, inset=0.12)
        p = tf.paragraphs[0]
        _add_runs(p, "우리에게 무엇인가  ", 10, accent, bold=True)
        _add_runs(p, s["sowhat"], 11)

    # 발표자 노트: 줄글판
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text
    return slide


def build_pptx(topic, per_category=5, link_base="../../docs/papers", since=2025):
    summaries, classification, timeline, insights, index = B.load_corpus(topic)
    slides, stats = B.build_slides(topic, summaries, classification, timeline, insights, index,
                                   per_category=per_category, link_base=link_base, since=since)
    essay_slides, _ = E.build(topic, per_category=per_category, link_base=link_base, since=since)
    notes_by_no = {}
    for es in essay_slides:
        parts = [es.get("lead", "")] + list(es.get("body", [])) + [es.get("close", "")]
        notes_by_no[es["no"]] = "\n\n".join(re.sub(r"\*\*(.+?)\*\*", r"\1", t) for t in parts if t)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    accent = _rgb(B.ACCENT.get(topic, "#D63423"))
    for s in slides:
        add_slide(prs, s, accent, notes_by_no.get(s["no"], ""))
    return prs, slides, stats


def main():
    ap = argparse.ArgumentParser(description="토픽 코퍼스 → 발표 슬라이드 PPTX")
    ap.add_argument("--topic", default="ai4s")
    ap.add_argument("--per-category", type=int, default=5)
    ap.add_argument("--since", type=int, default=2025)
    ap.add_argument("--link-base", default="../../docs/papers",
                    help="레퍼런스 링크 베이스(PPTX 위치 기준 상대경로 또는 절대 URL)")
    ap.add_argument("--out-dir", default=None, help="출력 루트 (기본 reports/)")
    args = ap.parse_args()

    try:
        prs, slides, stats = build_pptx(args.topic, args.per_category, args.link_base, args.since)
    except B.TaxonomyMismatch as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        return 2
    root = Path(args.out_dir) if args.out_dir else REPORTS
    out = root / "build" / f"{args.topic}_slides_{len(slides)}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    n_refs = len({r["slug"] for s in slides for r in s.get("refs", [])})
    print(f"[OK] PPTX 슬라이드 {len(slides)}장 · 레퍼런스 {n_refs}편 · 발표자 노트 {sum(1 for s in prs.slides if s.has_notes_slide)}장")
    print(f"  PPTX : {out}")
    return 0


if __name__ == "__main__":
    from _env_guard import force_py312
    force_py312()
    sys.exit(main())
