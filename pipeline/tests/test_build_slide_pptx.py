"""build_slide_pptx: 덱 슬라이드 → PPTX 구조 검증 (합성 슬라이드 + 라이브 코퍼스 스모크)."""
import io
import re
import unittest

from pptx import Presentation
from pptx.util import Emu

from pipeline import build_slide_deck as B
from pipeline import build_slide_pptx as P


def _texts(slide):
    out = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            out.append(shp.text_frame.text)
        if shp.has_table:
            for r in shp.table.rows:
                out.extend(c.text for c in r.cells)
    return "\n".join(out)


class SyntheticSlideTests(unittest.TestCase):
    SLIDE = {
        "no": 7, "kind": "서브카테고리", "part": "제1부 · 테스트", "title_ko": "제목",
        "title_en": "Title", "badges": ["12편", "2024–2026", "가속"],
        "cat_meta": "테스트 고유 배정 100편", "context": "파트 도입 문장",
        "headline": "헤드라인 **굵게** 포함",
        "points": ["2026 FooBench가 **무엇**을 쟀다.", "2025 Bar-2가 나왔다."],
        "tools": ["FooBench", "Bar-2"],
        "refs": [{"slug": "1_a", "title": "FooBench: a benchmark", "author": "A 외", "date": "2026",
                  "citations": 1234, "url": "../../docs/papers/1_a/index.html"}],
        "sowhat": "시사점",
        "extra": {"관전 포인트": ["하나", "둘"]},
    }

    def _one(self, slide_dict, notes="노트"):
        prs = Presentation()
        prs.slide_width, prs.slide_height = P.W, P.H
        P.add_slide(prs, slide_dict, P._rgb("#D63423"), notes)
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return Presentation(buf).slides[0]

    def test_slide_carries_every_field(self):
        s = self._one(self.SLIDE)
        t = _texts(s)
        for needle in ("S07", "서브카테고리", "제1부 · 테스트", "제목", "Title", "[12편]", "테스트 고유 배정 100편",
                       "파트 도입 문장", "헤드라인 굵게 포함", "FooBench가 무엇을 쟀다", "Bar-2가 나왔다",
                       "관전 포인트", "하나", "대표 도구 · 시스템: FooBench · Bar-2",
                       "[1] FooBench: a benchmark", "인용 1,234", "우리에게 무엇인가", "시사점"):
            self.assertIn(needle, t, needle)
        self.assertNotIn("**", t, "마크다운 굵게 표식이 그대로 남으면 안 된다")
        self.assertEqual(s.notes_slide.notes_text_frame.text, "노트")

    def test_reference_is_hyperlinked(self):
        s = self._one(self.SLIDE)
        links = [r.hyperlink.address for shp in s.shapes if shp.has_text_frame
                 for p in shp.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
        self.assertEqual(links, ["../../docs/papers/1_a/index.html"])

    def test_bold_markdown_becomes_bold_runs(self):
        s = self._one(self.SLIDE)
        bold = [r.text for shp in s.shapes if shp.has_text_frame
                for p in shp.text_frame.paragraphs for r in p.runs if r.font.bold and r.text in ("굵게", "무엇")]
        self.assertEqual(sorted(bold), ["굵게", "무엇"])

    def test_table_slide_uses_plain_grid_and_fits(self):
        slide = dict(self.SLIDE, no=4, table={"head": ["대분류", "고유 배정", "중복 포함"],
                                              "rows": [[f"카테고리 {i}", str(100 - i), str(300 - i)]
                                                       for i in range(10)]})
        s = self._one(slide)
        tbl = next(shp for shp in s.shapes if shp.has_table)
        self.assertEqual(len(tbl.table.rows), 11)
        self.assertEqual(tbl.table._tbl.tblPr.get("firstRow"), "0")
        self.assertEqual(tbl.table._tbl.tblPr.get("bandRow"), "0")
        self.assertLessEqual(tbl.top + tbl.height, P.H - P.MARGIN, "표가 슬라이드 아래로 넘친다")
        self.assertIn("카테고리 9", _texts(s))

    def test_shapes_stay_inside_slide(self):
        s = self._one(self.SLIDE)
        for shp in s.shapes:
            self.assertGreaterEqual(shp.left, 0)
            self.assertGreaterEqual(shp.top, 0)
            self.assertLessEqual(shp.left + shp.width, P.W + Emu(1))
            self.assertLessEqual(shp.top + shp.height, P.H + Emu(1), shp.name)

    def test_korean_runs_declare_gothic_east_asian_font(self):
        s = self._one(self.SLIDE)
        from pptx.oxml.ns import qn
        eas = {r._r.rPr.find(qn("a:ea")).get("typeface") for shp in s.shapes if shp.has_text_frame
               for p in shp.text_frame.paragraphs for r in p.runs if r._r.rPr is not None
               and r._r.rPr.find(qn("a:ea")) is not None}
        self.assertEqual(eas, {P.FONT_EA})


class LiveCorpusSmokeTests(unittest.TestCase):
    """docs/ai4s 가 있을 때 PPTX 를 끝까지 만들고 덱(HTML/MD)과 장수·노트·링크가 맞는지 본다."""

    def test_pptx_matches_deck(self):
        from pathlib import Path
        root = Path(__file__).resolve().parents[2] / "docs"
        if not all((root / "ai4s" / f).exists() for f in ("_new_classification.json", "_category_summaries.json",
                                                           "_timeline_narrative.json")):
            self.skipTest("docs/ai4s 코퍼스 없음")
        try:
            prs, slides, stats = P.build_pptx("ai4s")
        except B.TaxonomyMismatch as exc:
            self.skipTest(str(exc))
        self.assertEqual(len(prs.slides), len(slides))
        self.assertTrue(all(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides),
                        "모든 장에 줄글 발표자 노트가 있어야 한다")
        body = [s for s in slides if s.get("kind") == "서브카테고리"]
        self.assertEqual(len(body), stats["n_body"])
        for s, ps in zip(slides, prs.slides):
            t = _texts(ps)
            self.assertIsNone(re.search(r"\{[a-z_]+(?::[^{}]+)?\}", t), f"S{s['no']:02d}: 미치환 플레이스홀더")
            if s.get("refs"):
                links = [r.hyperlink.address for shp in ps.shapes if shp.has_text_frame
                         for p in shp.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
                self.assertTrue(links, f"S{s['no']:02d}: 레퍼런스 링크 없음")
