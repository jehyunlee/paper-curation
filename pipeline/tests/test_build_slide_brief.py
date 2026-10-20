"""build_slide_brief: 정책 브리핑 PPTX — 텍스트 정리 규칙·도판 선택·기하 정수화, 라이브 코퍼스 스모크."""
import io
import re
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

from pipeline import build_slide_brief as Bf
from pipeline import build_slide_deck as B


class TextRuleTests(unittest.TestCase):
    def test_ko_names_localises_categories_and_drops_slug_refs(self):
        t = ("Autonomous AI Scientific Discovery 카테고리와 Formal Mathematical Theorem Proving에서 "
             "예측 10900/10714, 모델 10566 (10909, 11016) Nature 624, 86 (2023) 44% 150편 2,709편")
        out = Bf.ko_names(t)
        self.assertIn("자율 AI 과학 발견와", out)
        self.assertIn("형식 수학 정리 증명에서", out)
        self.assertNotIn("10900", out)
        self.assertNotIn("10909", out)
        self.assertIn("(2023)", out, "연도는 근거 번호가 아니다")
        self.assertIn("44%", out)
        self.assertIn("150편", out)
        self.assertIn("2,709편", out)

    def test_strip_srefs_maps_slide_numbers_to_names(self):
        names = {17: "LLM 형식 정리 증명", 13: "에이전틱 신약 발굴"}
        self.assertEqual(Bf.strip_srefs("검증기 결합 (S17·S13) 이후", names), "검증기 결합 이후")
        self.assertEqual(Bf.strip_srefs("S17 참고", names), "'LLM 형식 정리 증명' 참고")

    def test_shorten_prefers_sentence_boundary(self):
        t = "첫 문장이다. 둘째 문장은 조금 더 길다, 그래서 잘린다. 셋째 문장."
        out = Bf.shorten(t, 30)
        self.assertTrue(out.endswith("다.") or out.endswith("다") or out.endswith("…"), out)
        self.assertLessEqual(len(out), 31)
        self.assertEqual(Bf.shorten("짧다", 30), "짧다")

    def test_shorten_cuts_at_clause_connective_not_noun_phrase_comma(self):
        t = "보정이 관측되는 데이터와 부합하는지를, 이론을 이용해 확인하는 검정을 제안하며, 상보적 문제를 다룬다."
        self.assertGreater(len(t), 45)
        out = Bf.shorten(t, 45)
        self.assertFalse(out.endswith("부합하는지를"), out)   # 명사구 쉼표에서 자르면 조사로 끝난다
        self.assertTrue(out.endswith("제안하며") or out.endswith("…"), out)

    def test_nominalize_converts_declarative_endings(self):
        cases = {
            "구조를 보여 준다.": "구조를 보여 줌.",
            "성능이 높다. 재현은 어렵다.": "성능이 높음. 재현은 어려움.",
            "GPT-5로 평가했다": "GPT-5로 평가했음",
            "이것은 중요한 문제이다.": "이것은 중요한 문제임.",
            "한계는 없다.": "한계는 없음.",
            "결과를 받는다": "결과를 받음",
            "값이 크다": "값이 큼",
            "자동화하는 멀티에이전트 시스템이며": "자동화하는 멀티에이전트 시스템임",
            "왔음을 지적하고": "왔음을 지적함",
            "응용에 확산되고 있어": "응용에 확산되고 있음",
            "결과가 좋지만": "결과가 좋음",
            "벤치마크 3종에서 SOTA": "벤치마크 3종에서 SOTA",
            "실험을 수행함": "실험을 수행함",
            "잘린 문장…": "잘린 문장…",
        }
        for src, want in cases.items():
            self.assertEqual(Bf.nominalize(src), want, src)

    def test_text_lines_counts_wrapping_and_newlines(self):
        self.assertEqual(Bf.text_lines("짧다", Inches(3), 10.5), 1)
        self.assertGreaterEqual(Bf.text_lines("가" * 60, Inches(2), 10.5), 3)
        self.assertEqual(Bf.text_lines("가\n나", Inches(3), 10.5), 2)

    def test_body_font_scale_leaves_titles_alone(self):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _, tf = Bf.textbox(s, 0, 0, Inches(3), Inches(1))
        Bf.add_runs(tf.paragraphs[0], "본문", 10, Bf.INK)
        Bf.add_runs(tf.paragraphs[0], "제목", 24, Bf.INK, scale=False)
        sizes = [r.font.size.pt for r in tf.paragraphs[0].runs]
        self.assertAlmostEqual(sizes[0], 10 * Bf.FS, places=1)
        self.assertEqual(sizes[1], 24)

    def test_source_links_prefers_doi_then_url(self):
        import os
        import sqlite3
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            db = Path(d) / "b.sqlite3"
            con = sqlite3.connect(db)
            con.execute("CREATE TABLE papers (slug TEXT, doi TEXT, url TEXT)")
            con.executemany("INSERT INTO papers VALUES (?,?,?)", [
                ("a", "10.1000/xyz", "https://arxiv.org/abs/1"),
                ("b", "", "https://openreview.net/forum?id=Q"),
                ("c", None, None),
            ])
            con.commit()
            con.close()
            old = os.environ.get("PAPER_CURATION_BIBLIO_DB")
            os.environ["PAPER_CURATION_BIBLIO_DB"] = str(db)
            try:
                links = Bf.source_links({"a", "b", "c", "zzz"})
            finally:
                if old is None:
                    del os.environ["PAPER_CURATION_BIBLIO_DB"]
                else:
                    os.environ["PAPER_CURATION_BIBLIO_DB"] = old
        self.assertEqual(links["a"], ("doi.org/10.1000/xyz", "https://doi.org/10.1000/xyz"))
        self.assertEqual(links["b"][1], "https://openreview.net/forum?id=Q")
        self.assertNotIn("https://", links["b"][0])
        self.assertNotIn("c", links)
        self.assertNotIn("zzz", links)

    def test_strip_year_removes_leading_year_or_background_label(self):
        self.assertEqual(Bf.strip_year("2026 FooBench가 쟀다."), "FooBench가 쟀다.")
        self.assertEqual(Bf.strip_year("2018–2021 SCANPY류"), "SCANPY류")
        self.assertEqual(Bf.strip_year("배경: 과거 흐름"), "과거 흐름")

    def test_E_rounds_to_integer_emu(self):
        v = Bf.E(Bf.INNER_W * 0.6)
        self.assertIsInstance(v, int)
        self.assertEqual(v, int(round(Bf.INNER_W * 0.6)))


class FigureTests(unittest.TestCase):
    def test_good_figure_filters_small_or_extreme_aspect(self):
        figs = [{"w": 300, "h": 300, "n": 1}, {"w": 1400, "h": 300, "n": 2}, {"w": 1200, "h": 800, "n": 3}]
        self.assertEqual(Bf.good_figure(figs)["n"], 3)
        self.assertIsNone(Bf.good_figure(figs[:2]))

    def test_autocrop_trims_white_margins_but_keeps_tiny_content(self):
        im = Image.new("RGB", (400, 300), "white")
        for x in range(80, 330):
            for y in range(50, 250):
                im.putpixel((x, y), (30, 60, 200))
        out = Bf.autocrop(im)
        self.assertLess(out.width, 300)
        self.assertLess(out.height, 240)
        dot = Image.new("RGB", (400, 300), "white")
        dot.putpixel((5, 5), (0, 0, 0))
        self.assertEqual(Bf.autocrop(dot).size, (400, 300), "내용이 거의 없으면 원본 유지")


class GeometryTests(unittest.TestCase):
    def _slide(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Bf.W, Bf.H
        return prs, prs.slides.add_slide(prs.slide_layouts[6])

    def test_chart_frame_has_integer_extent(self):
        prs, s = self._slide()
        b = Bf.Brief.__new__(Bf.Brief)
        b.prs = prs
        b.bar_chart(s, Bf.M, Bf.TOP, Bf.INNER_W * 0.6, Inches(3.3), ["가", "나"], [1, 2], Bf.ACCENT)
        buf = io.BytesIO()
        prs.save(buf)
        xml = Presentation(io.BytesIO(buf.getvalue())).slides[0]._element.xml
        self.assertIsNone(re.search(r'(?:cx|cy|x|y)="\d+\.\d+"', xml), "float EMU 가 남으면 PowerPoint 가 차트를 버린다")

    def test_fit_picture_keeps_aspect_and_stays_inside_box(self):
        prs, s = self._slide()
        tmp = Path("/tmp/_brief_test.png")
        Image.new("RGB", (1000, 400), "red").save(tmp)
        box = (Inches(1), Inches(1), Inches(4), Inches(4))
        pic = Bf.fit_picture(s, tmp, *box)
        self.assertAlmostEqual(pic.width / pic.height, 2.5, places=1)
        self.assertGreaterEqual(pic.left, box[0])
        self.assertLessEqual(pic.left + pic.width, box[0] + box[2] + Emu(1))
        self.assertLessEqual(pic.top + pic.height, box[1] + box[3] + Emu(1))


class LiveCorpusSmokeTests(unittest.TestCase):
    def test_brief_builds_with_figures_notes_and_no_placeholders(self):
        root = Path(__file__).resolve().parents[2] / "docs" / "ai4s"
        if not all((root / f).exists() for f in ("_new_classification.json", "_category_summaries.json",
                                                "_timeline_narrative.json", "_insights.json", "_umap_coords.json")):
            self.skipTest("docs/ai4s 코퍼스 없음")
        try:
            prs, brief = Bf.build_brief("ai4s")
        except B.TaxonomyMismatch as exc:
            self.skipTest(str(exc))
        n_cats = len(brief.corpus["cats"])
        self.assertGreaterEqual(len(prs.slides), 3 + 5 + 2 * n_cats + 10)
        n_pics = sum(1 for s in prs.slides for shp in s.shapes if shp.shape_type == 13)  # PICTURE
        self.assertGreaterEqual(n_pics, 2 * n_cats + 1, "대분류마다 도판 2장 이상 + 지형도")
        n_charts = sum(1 for s in prs.slides for shp in s.shapes if shp.has_chart)
        self.assertGreaterEqual(n_charts, n_cats + 2)
        for s in prs.slides:
            slide_text = "\n".join(shp.text_frame.text for shp in s.shapes if shp.has_text_frame)
            is_sources = "도판 출처" in slide_text   # 부록 C 는 이 덱 자신의 S번호를 쓴다
            for shp in s.shapes:
                self.assertLessEqual(shp.left + shp.width, Bf.W + Emu(1), shp.name)
                self.assertLessEqual(shp.top + shp.height, Bf.H + Emu(1), shp.name)
                if shp.has_text_frame:
                    self.assertIsNone(re.search(r"\{[a-z_]+(?::[^{}]+)?\}", shp.text_frame.text),
                                      "미치환 플레이스홀더")
                    if not is_sources:
                        self.assertNotRegex(shp.text_frame.text, r"\bS\d{2}\b", "58장 덱의 슬라이드 번호가 새면 안 된다")
        self.assertTrue(any(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides))
        # 문체 계약: 슬라이드 본문에 서술형 종결('~다.')·메타 문구·저자명 출처가 없고, 출처는 DOI/URL 링크
        verb = re.compile(r"(?:[가-힣]다|[어아해예지네]요|죠)\s*[.!]?\s*$")
        meta = re.compile(r"도판은|읽는 법|핵심.{0,3}으로 고른|출처:|[A-Z][a-z]+ [A-Z][a-z]+ 외")
        n_links = 0
        for s in prs.slides:
            for shp in s.shapes:
                if not shp.has_text_frame:
                    continue
                for para in shp.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs)
                    n_links += sum(1 for r in para.runs if r.hyperlink and r.hyperlink.address)
                    for sent in re.split(r"(?<=[.!])\s+|\n", text):
                        self.assertIsNone(verb.search(sent), f"서술형 종결: {sent[-50:]}")
                        self.assertIsNone(meta.search(sent), f"메타/저자 출처 문구: {sent[:60]}")
                    if "Fig." in text:
                        self.assertTrue(any(r.hyperlink and r.hyperlink.address for r in para.runs), text)
        self.assertGreaterEqual(n_links, len(brief.sources))
        self.assertEqual(len({(p["slug"], f["n"]) for _, p, f in brief.sources}),
                         len({(p["slug"], f["n"]) for _, p, f in brief.sources}))
        xml = "".join(s._element.xml for s in prs.slides)
        self.assertIsNone(re.search(r'(?:cx|cy)="\d+\.\d+"', xml))
